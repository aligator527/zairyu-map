"""Build one unified, sorted long-format dataset from the downloaded
在留外国人統計テーブルデータ（国籍・地域別 在留資格別 市区町村別） files (t2, since 2023-12).

Output columns:

    period, year, month,
    prefecture_code, prefecture, designated_city_code, designated_city,
    municipality_code, municipality,
    region_code, region, nationality_code, nationality,
    status_code, status,
    sex_code, sex,
    age, age_label, age_group_code, age_group,
    count

What differs between periods and how it is handled:
  * 2023-12, 2024-06: municipality x nationality x status only. sex/age are NA.
    Municipalities with <= 10 foreigners are lumped into one nationwide row
    municipality_code 99999 "その他" (prefecture_code 99).
  * 2024-12+: sex and age are added. The data lives in a Power Pivot model,
    not on a worksheet, and is read with pbixray. Values suppressed by the
    Ministry ("-") become "秘匿" in the text columns and NA in the code columns:
    municipalities with < 5,000 foreigners have age/sex suppressed, those with
    <= 10 have everything but the count suppressed. In these files code 99999
    means 未定・不詳; it is mapped to 48000 / prefecture 48, as in 2023-12/2024-06.
  * 2023-12 has names without codes: nationality/status/prefecture codes are
    filled in from the other periods (and the t1 dataset for nationalities).
  * designated_city (政令指定都市) exists only from 2024-12; for older periods it
    is filled in from the municipality code.

Usage:
    python build_municipal.py      # -> data/zairyu_gaikokujin_municipal.{parquet,csv.gz}
"""

from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path

import openpyxl
import pandas as pd

from build_dataset import (
    INTERIM_DIR, RAW_DIR, REGIONS, SEX,
    age_labels, build_nationality_codes, code_name, nationality_name, parse_age,
)

OUT_BASE = Path("data/zairyu_gaikokujin_municipal")
SUPPRESSED = "秘匿"

log = logging.getLogger("build_municipal")


# ---------------------------------------------------------------- reading


def _patch_pbixray() -> None:
    """pbixray skips table files whose name starts with 'H' or 'R' (it means to
    skip only the internal 'H$'/'R$' tables); the 2024-12+ models are named
    'R712 テーブルデータ用BD' etc., so they would decode as empty."""
    from pbixray.meta import xml_source as xs

    rx = re.compile(r"^(?![HR]\$)([^$]+?)\.(\d+)\.tbl\.xml$")

    def _extract_tbl_metadata(self):
        self._load_xml_collection(
            rx, lambda s: xs.XMObjectDocument.from_xml_string(s).root_object,
            self._tbl_objects, "table",
        )

    xs.XmlMetadataSource._extract_tbl_metadata = _extract_tbl_metadata


def read_sheet(path: Path) -> pd.DataFrame | None:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if header and "在留外国人数" in header and "市区町村コード" in header and header[-1] == "在留外国人数":
                df = pd.DataFrame((r for r in rows if any(v is not None for v in r)),
                                  columns=[str(h) for h in header])
                log.info("%s: sheet %r, %d rows", path.name, ws.title, len(df))
                return df
    finally:
        wb.close()
    return None


def read_power_pivot(path: Path) -> pd.DataFrame:
    _patch_pbixray()
    from pbixray import PBIXRay

    model = PBIXRay(str(path))
    try:
        for table in model.tables:
            df = model.get_table(table)
            if "在留外国人数" in df.columns and "市区町村コード" in df.columns:
                if df.empty:
                    raise ValueError(f"{path}: Power Pivot table {table!r} decoded as empty")
                log.info("%s: Power Pivot table %r, %d rows", path.name, table, len(df))
                return df.drop(columns=["__XL_RowNumber"], errors="ignore")
    finally:
        model.close()
    raise ValueError(f"{path}: no municipal data found (neither sheet nor Power Pivot)")


def load_raw(period: str, path: Path, use_cache: bool) -> pd.DataFrame:
    cache = INTERIM_DIR / f"{period}_t2.parquet"
    if use_cache and cache.exists() and cache.stat().st_mtime >= path.stat().st_mtime:
        return pd.read_parquet(cache)
    df = read_sheet(path)
    if df is None:
        df = read_power_pivot(path)
    df = df.astype({c: "string" for c in df.columns if c != "在留外国人数"})
    INTERIM_DIR.mkdir(parents=True, exist_ok=True)
    df.to_parquet(cache, index=False)
    return df


# ---------------------------------------------------------------- lookups


def _pairs(df: pd.DataFrame, col: str) -> dict[str, str]:
    """name -> code from '35：永住者'-style values."""
    if col not in df.columns:
        return {}
    code, name = code_name(df[col].drop_duplicates())
    return dict(zip(name.dropna(), code.dropna()))


def build_lookups(raw: dict[str, pd.DataFrame]) -> dict[str, dict[str, str]]:
    status, pref, city = {}, {}, {}
    nat_values = {p: set(df["国籍・地域"].dropna()) for p, df in raw.items()}
    # The prefecture-level dataset knows every nationality code seen since 2016.
    for f in sorted(INTERIM_DIR.glob("*_t1.parquet")):
        nat_values[f"t1:{f.name[:7]}"] = set(pd.read_parquet(f, columns=["国籍・地域"])["国籍・地域"].dropna())
    for p in sorted(raw):
        df = raw[p]
        status |= _pairs(df, "在留資格")
        pref |= _pairs(df, "都道府県")
        if "政令指定都市" in df.columns:
            c = df[["市区町村コード", "政令指定都市"]].dropna().drop_duplicates()
            city |= dict(zip(c["市区町村コード"], c["政令指定都市"]))
    return {"nat": build_nationality_codes(nat_values), "status": status, "pref": pref, "city": city}


# ---------------------------------------------------------------- normalising


def _split_or_lookup(s: pd.Series, lookup: dict[str, str]) -> tuple[pd.Series, pd.Series]:
    """'35：永住者' -> ('35','永住者'); bare '永住者' -> code via lookup; '-' -> (NA, 秘匿)."""
    s = s.str.strip()
    code, name = code_name(s)
    bare = code.isna() & (s != "-")
    name = name.mask(bare, s).mask(s == "-", SUPPRESSED)
    code = code.mask(bare, s.map(lookup))
    return code, name


def normalise(period: str, df: pd.DataFrame, lk: dict[str, dict[str, str]]) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    year, month = period.split("-")
    out["period"] = period
    out["year"] = int(year)
    out["month"] = int(month)

    muni_code = df["市区町村コード"].str.strip()
    muni = df["市区町村"].str.strip()
    # 2024-12+: 99999 = 未定・不詳 (older files use 48000 for that, 99999 for その他).
    unknown_place = (muni_code == "99999") & (muni == "未定・不詳")
    muni_code = muni_code.mask(unknown_place, "48000")
    pref_raw = df["都道府県"].str.strip().mask(unknown_place, "48：未定・不詳")
    out["prefecture_code"], out["prefecture"] = _split_or_lookup(pref_raw, lk["pref"])
    out["prefecture_code"] = out["prefecture_code"].fillna(muni_code.str[:2])
    city = df["政令指定都市"] if "政令指定都市" in df.columns else pd.Series(pd.NA, index=df.index, dtype="string")
    city = city.fillna(muni_code.map(lk["city"]))
    out["designated_city_code"], out["designated_city"] = code_name(city.astype("string"))
    out["municipality_code"] = muni_code
    out["municipality"] = muni

    nat = df["国籍・地域"].str.strip()
    names = nat.map(lambda v: SUPPRESSED if v == "-" else nationality_name(v))
    out["nationality_code"] = names.map(lk["nat"])
    out["region_code"] = out["nationality_code"].str[:2]
    out["region"] = out["region_code"].map(REGIONS)
    out["nationality"] = names
    out["status_code"], out["status"] = _split_or_lookup(df["在留資格"], lk["status"])

    if "性別" in df.columns:
        sex = df["性別"].str.strip()
        sex_code, _ = code_name(sex)
        out["sex_code"] = pd.to_numeric(sex_code).astype("Int64")
        out["sex"] = sex_code.str.lstrip("0").map(SEX).mask(sex == "-", SUPPRESSED)
        age_raw = df["年齢"].str.strip()
        out["age"] = parse_age(age_raw)
        labels = age_labels(out["age"])
        hidden = age_raw == "-"
        labels.loc[hidden, ["age_label", "age_group"]] = SUPPRESSED
        out = out.join(labels)
    else:  # not published for this period
        out["sex_code"] = pd.array([pd.NA] * len(df), dtype="Int64")
        out["age"] = pd.array([pd.NA] * len(df), dtype="Int64")
        out["age_group_code"] = pd.array([pd.NA] * len(df), dtype="Int64")
        for c in ("sex", "age_label", "age_group"):
            out[c] = pd.Series(pd.NA, index=df.index, dtype="string")
    out["count"] = pd.to_numeric(df["在留外国人数"]).astype("int64")

    for col in ("prefecture_code", "municipality_code"):
        if out[col].isna().any():
            raise ValueError(f"{period}: missing {col}: {df.loc[out[col].isna()].head(3).to_dict('records')}")
    for col, name_col in (("nationality_code", "nationality"), ("status_code", "status")):
        miss = out[col].isna() & (out[name_col] != SUPPRESSED)
        if miss.any():
            raise ValueError(f"{period}: no code for {name_col}: {sorted(out.loc[miss, name_col].unique())}")

    keys = [c for c in out.columns if c != "count"]
    return out.groupby(keys, dropna=False, sort=False, observed=True, as_index=False)["count"].sum()


SORT_KEYS = ["period", "prefecture_code", "municipality_code", "nationality_code", "status_code", "age", "sex_code"]
COLUMNS = [
    "period", "year", "month",
    "prefecture_code", "prefecture", "designated_city_code", "designated_city",
    "municipality_code", "municipality",
    "region_code", "region", "nationality_code", "nationality",
    "status_code", "status",
    "sex_code", "sex", "age", "age_label", "age_group_code", "age_group",
    "count",
]
CATEGORY_COLS = [
    "period", "prefecture_code", "prefecture", "designated_city_code", "designated_city",
    "municipality_code", "municipality", "region_code", "region", "nationality_code", "nationality",
    "status_code", "status", "sex", "age_label", "age_group",
]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--raw", type=Path, default=RAW_DIR)
    ap.add_argument("--out", type=Path, default=OUT_BASE, help="output path without extension")
    ap.add_argument("--no-csv", action="store_true")
    ap.add_argument("--no-cache", action="store_true", help="re-read xlsx even if data/interim has it")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    files = sorted(args.raw.glob("????-??_t2.xlsx"))
    if not files:
        raise SystemExit(f"no *_t2.xlsx in {args.raw}; run estat_crawler.py first")

    raw = {f.name[:7]: load_raw(f.name[:7], f, not args.no_cache) for f in files}
    lookups = build_lookups(raw)
    data = pd.concat([normalise(p, df, lookups) for p, df in raw.items()], ignore_index=True)[COLUMNS]
    data = data.sort_values(SORT_KEYS, na_position="last", kind="stable").reset_index(drop=True)
    data[CATEGORY_COLS] = data[CATEGORY_COLS].astype("category")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    data.to_parquet(args.out.with_suffix(".parquet"), index=False)
    log.info("wrote %s (%d rows)", args.out.with_suffix(".parquet"), len(data))
    if not args.no_csv:
        data.to_csv(args.out.with_suffix(".csv.gz"), index=False, encoding="utf-8")
        log.info("wrote %s", args.out.with_suffix(".csv.gz"))

    summary = data.groupby("period", observed=True).agg(
        total=("count", "sum"), municipalities=("municipality_code", "nunique"))
    summary["total"] = summary["total"].map("{:,}".format)
    print("\nTotal 在留外国人 by period:")
    print(summary.to_string())


if __name__ == "__main__":
    main()
